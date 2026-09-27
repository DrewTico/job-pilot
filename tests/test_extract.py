"""Resume extraction layouts, using only synthetic documents."""

import pytest
from docx import Document

from job_agent.tailor.career_facts import load_career_facts
from job_agent.tailor.extract import build_career_facts, extract_career_facts, write_career_facts


def resume(tmp_path, paragraphs):
    doc = Document()
    for text, bold in paragraphs:
        doc.add_paragraph().add_run(text).bold = bold
    path = tmp_path / 'synthetic.docx'
    doc.save(path)
    return path


def test_legacy_layout(tmp_path):
    path = resume(tmp_path, [
        ('Casey Example', True), ('Data Engineer', False),
        ('TECHNICAL SKILLS', True), ('Languages', True), ('Python, SQL', False),
        ('EDUCATION', True), ('Example College', False),
        ('WORK EXPERIENCE', True),
        ("Example Labs's | Remote | Data Engineer  Jan 2020 - Present", True),
        ('Reporting platform.', False), ('Key Responsibilities', True),
        ('Reduced latency by 25%; served 200 users.', False),
        ('Environment: Python, Service (A, B), SQL', False),
        ('Sample Works | Sample City | Analyst  Feb 2018 – Dec 2019', True),
        ('Analytics.', False), ('Key Responsibilities', True), ('Built reports.', False),
    ])
    facts = extract_career_facts(path)
    assert facts['name'] == 'Casey Example'
    assert facts['role'] == 'Data Engineer'
    assert facts['education'] == ['Example College']
    assert facts['skills_inventory'] == {'Languages': ['Python, SQL']}
    assert 'projects' not in facts
    assert facts['employers'][0] == {
        'company': 'Example Labs', 'location': 'Remote', 'title': 'Data Engineer',
        'duration': 'Jan 2020 - Present', 'project_description': 'Reporting platform.',
        'real_bullets': ['Reduced latency by 25%; served 200 users.'],
        'real_metrics': ['Reduced latency by 25%', 'served 200 users'],
        'real_skills': ['Python', 'Service (A, B)', 'SQL'],
    }
    assert facts['employers'][1]['real_bullets'] == ['Built reports.']


@pytest.mark.parametrize('skills_heading', ['SKILLS', 'TECHNICAL SKILLS'])
@pytest.mark.parametrize('work_heading', ['EXPERIENCE', 'WORK EXPERIENCE'])
def test_split_entries_inline_skills_and_projects(tmp_path, skills_heading, work_heading):
    path = resume(tmp_path, [
        ('Casey Example', True), ('Developer', False),
        ('EDUCATION', True), ('Example College\tMay 2017', True),
        ('B.A. Computer Science', False),
        (skills_heading, True),
        ('• Languages: Python, SQL', False),
        ('• AI/ML: Classification, Evaluation', False),
        ('• Cloud & Software Development: Containers, API design', False),
        ('• Tools & Frameworks: Widget (A, B), ExampleKit', False),
        (work_heading, True),
        ('Example Labs\tJanuary 2021 - Present', True),
        ('Software Developer\tSample City, ZZ', False),
        ('Built a queue serving 400 users.', False), ('Maintained tests.', False),
        ('Sample Works\tFeb. 2018 – Dec. 2020', True),
        ('Platform Engineer\tRemote', False), ('Improved monitoring.', False),
        ('PROJECTS', True), ('Paper Kite - Task planner\t2022', True),
        ('Built a scheduling prototype.', False),
        ('Map Lantern\t2019 - 2020', True), ('Added route previews.', False),
    ])
    facts = build_career_facts(path, email='casey@example.invalid', phone='')
    assert facts['education'] == ['Example College\tMay 2017', 'B.A. Computer Science']
    assert facts['skills_inventory'] == {
        'Languages': ['Python', 'SQL'], 'AI/ML': ['Classification', 'Evaluation'],
        'Cloud & Software Development': ['Containers', 'API design'],
        'Tools & Frameworks': ['Widget (A, B)', 'ExampleKit'],
    }
    assert len(facts['employers']) == 2
    assert facts['employers'][0] == {
        'company': 'Example Labs', 'title': 'Software Developer',
        'location': 'Sample City, ZZ', 'duration': 'January 2021 - Present',
        'project_description': '',
        'real_bullets': ['Built a queue serving 400 users.', 'Maintained tests.'],
        'real_metrics': ['Built a queue serving 400 users'], 'real_skills': [],
    }
    assert facts['employers'][1]['real_bullets'] == ['Improved monitoring.']
    assert facts['projects'] == [
        {'header': 'Paper Kite - Task planner\t2022',
         'real_bullets': ['Built a scheduling prototype.']},
        {'header': 'Map Lantern\t2019 - 2020', 'real_bullets': ['Added route previews.']},
    ]
    loaded = load_career_facts(write_career_facts(facts, tmp_path / 'synthetic.yaml'))
    assert loaded.model_dump(mode='json')['projects'] == facts['projects']
    assert isinstance(loaded.projects, tuple)
    assert isinstance(loaded.projects[0].real_bullets, tuple)
    with pytest.raises(ValueError):
        loaded.projects[0].header = 'Changed'


@pytest.mark.parametrize('contact', [
    'casey@example.invalid',
    '(202) 555-0147',
    '+1 202-555-0147',
    '2025550147',
    'https://www.linkedin.com/in/casey-example',
    'github.com/casey-example',
    'Sample City | Portfolio | Contact',
    'Sample City | casey@example.invalid | (202) 555-0147',
])
def test_contact_line_is_not_role(tmp_path, contact):
    path = resume(tmp_path, [
        ('Casey Example', True), (contact, False), ('SKILLS', True),
        ('Languages: Python', False),
    ])
    assert extract_career_facts(path)['role'] == ''


@pytest.mark.parametrize('role', ['Software Developer', 'Engineer II', 'Developer | Data Tools'])
def test_real_role_is_preserved(tmp_path, role):
    path = resume(tmp_path, [('Casey Example', True), (role, False)])
    assert extract_career_facts(path)['role'] == role


def test_section_heading_is_not_role(tmp_path):
    path = resume(tmp_path, [('Casey Example', True), ('SUMMARY', True)])
    assert extract_career_facts(path)['role'] == ''


def test_inline_skills_with_mixed_runs_and_no_education(tmp_path):
    doc = Document()
    doc.add_paragraph('Casey Example')
    doc.add_paragraph('Developer')
    doc.add_paragraph('SKILLS')
    p = doc.add_paragraph()
    p.add_run('• Languages: ').bold = True
    p.add_run('Python, SQL')
    doc.add_paragraph('PROJECTS')
    doc.add_paragraph().add_run('Practice App').bold = True
    doc.add_paragraph('SKILLS improved through practice.')
    path = tmp_path / 'synthetic.docx'
    doc.save(path)
    facts = extract_career_facts(path)
    assert facts['skills_inventory'] == {'Languages': ['Python', 'SQL']}
    assert facts['education'] == []
    assert facts['employers'] == []
    assert facts['projects'][0]['real_bullets'] == ['SKILLS improved through practice.']
